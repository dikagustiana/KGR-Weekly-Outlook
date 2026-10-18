#!/usr/bin/env python3
"""B-R91 (PRD v5.2-B, K70): SQL that imports the FS Management Report rows of given months into fs_statement_line.

  python3 -I scripts/fs_import.py                write docs/v5.2/kgr-os-data-fs.sql (August and September 2026)
  python3 -I scripts/fs_import.py --check        write nothing; exit 1 when the file on disk differs

Source: reference/09._Laporan_Keuangan_Per_Sep_2026_PT_KGR_v2.xlsx, sheets "PL - Management Report" and
"BS - Management Report", the month-end column of each month (first matching date in header row 10, columns D–L, as
scripts/build_fs_recon.py reads it). A row is imported when column B has a label and the month cell holds a number;
row 70 of the PL sheet ("Current year per BS", label in column C) is a cross-check of the workbook, not a statement row.
Amounts are numeric(20,2), rounded half up. The rows read here must equal the rows of fixtures/fs-2026-09/inputs.json
(A4) and the PRD check figures (revenue, GP, NIBT); the script stops otherwise. Nothing is adjusted (K62).

The SQL refuses to run before bundle B5 (schema_version() >= 5.2.4), inserts version 1 with locked_at set (Q-v5.2-66:
the import locks the rows; a correction is a new version), skips rows that are already there (re-runnable) and ends
with a check table whose every row must be true.
"""
import datetime as dt
import hashlib
import json
import os
import sys
from decimal import ROUND_HALF_UP, Decimal as D

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FS_XLSX = "reference/09._Laporan_Keuangan_Per_Sep_2026_PT_KGR_v2.xlsx"
OUT = "docs/v5.2/kgr-os-data-fs.sql"
MONTHS = [("2026-08-01", "2026-08-31"), ("2026-09-01", "2026-09-30")]
SHEETS = [("PL", "pl", "PL - Management Report"), ("BS", "bs", "BS - Management Report")]
# PRD v5.2-B B-R91 check figures, in Rp jt with one decimal (PL rows 12 Revenue, 15 GROSS PROFIT, 51 INCOME (LOSS) BEFORE TAX)
CHECK = {
    "2026-08-01": {12: D("8361.8"), 15: D("306.7"), 51: D("-26.5")},
    "2026-09-01": {12: D("5039.9"), 15: D("-475.4"), 51: D("-910.1")},
}


def q2(v):
    return D(repr(v) if isinstance(v, float) else str(v)).quantize(D("0.01"), rounding=ROUND_HALF_UP)


def sha256(rel):
    h = hashlib.sha256()
    with open(os.path.join(ROOT, rel), "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def lit(s):
    return "'" + s.replace("'", "''") + "'"


def read_rows():
    wb = openpyxl.load_workbook(os.path.join(ROOT, FS_XLSX), data_only=True)
    fx = json.load(open(os.path.join(ROOT, "fixtures/fs-2026-09/inputs.json"), encoding="utf-8"))
    rows = []
    for month, end in MONTHS:
        for st, key, sheet in SHEETS:
            ws = wb[sheet]
            col = next(c for c in range(4, 13)
                       if isinstance(ws.cell(10, c).value, (dt.date, dt.datetime)) and ws.cell(10, c).value.date().isoformat() == end)
            got = {}
            for r in range(11, ws.max_row + 1):
                label, v = ws.cell(r, 2).value, ws.cell(r, col).value
                if label is None or str(label).strip() == "" or isinstance(v, bool) or not isinstance(v, (int, float)):
                    continue
                got[r] = (str(label).strip(), q2(v))
            want = {x["row"]: (x["line"], q2(x["values"][end])) for x in fx[key] if x["values"].get(end) is not None}
            if got != want:
                diff = sorted(set(got.items()) ^ set(want.items()))
                sys.exit(f"{sheet} {end}: workbook rows differ from fixtures/fs-2026-09 (A4): {diff[:6]}")
            rows += [(month, st, sheet, r, line, amt) for r, (line, amt) in sorted(got.items())]
    for month, figs in CHECK.items():
        for r, jt in figs.items():
            amt = next(a for m, st, _, rr, _, a in rows if m == month and st == "PL" and rr == r)
            if (amt / D(1_000_000)).quantize(D("0.1"), rounding=ROUND_HALF_UP) != jt:
                sys.exit(f"PL b.{r} {month}: {amt} != Rp {jt} jt (PRD B-R91)")
    return rows


def render(rows, sha):
    name = os.path.basename(FS_XLSX)
    counts = {}
    for m, st, *_ in rows:
        counts[(m, st)] = counts.get((m, st), 0) + 1
    total = {(m, st): sum(a for mm, s, _, _, _, a in rows if (mm, s) == (m, st)) for (m, st) in counts}
    values = ",\n".join(
        f"  ({lit(m)}::date, {lit(st)}, {lit(sh)}, {r}, {lit(line)}, {amt})" for m, st, sh, r, line, amt in rows)
    checks = []
    for (m, st), n in sorted(counts.items()):
        checks.append(f"  ('{st} {m[:7]}: {n} baris versi 1 dari berkas ini, terkunci, Σ = {total[(m, st)]}',\n"
                      f"   (select count(*) = {n} and count(locked_at) = {n} and sum(amount) = {total[(m, st)]}\n"
                      f"      and bool_and(source_sha256 = {lit(sha)} and source_file = {lit(name)})\n"
                      f"    from public.fs_statement_line where month = {lit(m)} and statement = {lit(st)} and version = 1))")
    for m, figs in CHECK.items():
        for r, jt in figs.items():
            checks.append(f"  ('PL {m[:7]} b.{r} = Rp {jt} jt (PRD B-R91)',\n"
                          f"   (select round(amount / 1000000, 1) = {jt} from public.fs_statement_line\n"
                          f"    where month = {lit(m)} and statement = 'PL' and row_no = {r} and version = 1))")
    checks.append("  ('tidak ada versi lain untuk bulan-bulan ini',\n"
                  f"   (select count(*) = 0 from public.fs_statement_line where month in ({', '.join(lit(m) for m, _ in MONTHS)}) and version <> 1))")
    return f"""-- KGR-OS: impor FS acuan (PRD v5.2-B B-R91, K70) ke fs_statement_line, bulan {', '.join(m[:7] for m, _ in MONTHS)}.
-- Dibangkitkan oleh scripts/fs_import.py (jangan disunting tangan; jalankan ulang skripnya). Sumber: {FS_XLSX}
-- sha256 {sha}; sheet "PL - Management Report" dan "BS - Management Report", kolom akhir bulan (baris 10).
-- Baris = label di kolom B dan angka di kolom bulan; sama dengan fixtures/fs-2026-09/inputs.json (A4) dan angka cek PRD.
-- Versi 1, locked_at diisi (Q-v5.2-66). Bisa dijalankan ulang: baris yang sudah ada tidak disentuh. Tidak memuat escape
-- unicode. Hasil yang benar: tabel cek di akhir, kolom ada semuanya true.

do $$
begin
  if to_regprocedure('public.schema_version()') is null or to_regclass('public.fs_statement_line') is null
     or string_to_array(coalesce(public.schema_version(), '0'), '.')::int[] < array[5,2,4] then
    raise exception 'fs_statement_line belum ada (bundel B5, schema_version() >= 5.2.4); impor FS tidak dijalankan';
  end if;
end $$;

insert into public.fs_statement_line (month, statement, sheet, row_no, line, amount, source_file, source_sha256, version, locked_at)
select v.month, v.statement, v.sheet, v.row_no, v.line, v.amount, {lit(name)}, {lit(sha)}, 1, now()
from (values
{values}
) as v(month, statement, sheet, row_no, line, amount)
on conflict (month, statement, row_no, version) do nothing;

select cek, ada from (values
{(',' + chr(10)).join(checks)}
) as t(cek, ada);
"""


def main():
    sha = sha256(FS_XLSX)
    text = render(read_rows(), sha)
    path = os.path.join(ROOT, OUT)
    if "--check" in sys.argv:
        same = os.path.exists(path) and open(path, encoding="utf-8").read() == text
        print(f"{OUT}: {'sama' if same else 'BERBEDA'}")
        sys.exit(0 if same else 1)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
    print(f"{OUT}: {text.count(chr(10))} baris, sha256 berkas sumber {sha}")


if __name__ == "__main__":
    main()
