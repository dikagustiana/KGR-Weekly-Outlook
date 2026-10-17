#!/usr/bin/env python3
"""Fixtures fs-recon-2026-08 / fs-recon-2026-09 and docs/v5.2/fs-recon-fixture.md (PRD v5.2-A R68, T15, K62, O25–O27).

  python3 -I scripts/build_fs_recon.py          write fixtures/fs-recon-2026-0{8,9}/ and docs/v5.2/fs-recon-fixture.md
  python3 -I scripts/build_fs_recon.py --check  write nothing; exit 1 when a file on disk differs from what this would write

FS side: the Management Report of reference/09._Laporan_Keuangan_Per_Sep_2026_PT_KGR_v2.xlsx through
fixtures/fs-2026-09/inputs.json (`pl[].row`, `bs[].row`); every row used is read from the workbook again here (label and
value) and must agree. App side: the repo's fixture data (fixtures/week-2026-08-24: journal 1–26 Aug 2026 and the trial
balance of 31 Jul 2026), folded by THIS script with plain Decimal arithmetic. The engine (packages/engine/src/fs-app.ts)
computes the same numbers independently; `packages/engine/test/fs-recon.test.ts` compares the two and then compares
fsRecon against golden.json. Nothing is adjusted to make a row pass: where the fixture data is incomplete the golden holds
the expected difference with its cause.
"""
import datetime as dt
import hashlib
import json
import os
import sys
from decimal import ROUND_HALF_UP, Decimal as D

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FS_XLSX = "reference/09._Laporan_Keuangan_Per_Sep_2026_PT_KGR_v2.xlsx"
FS_SHA = "7e68aa89751cd8938bd5c9cd4f961ebbbf8eed56ac55cf3ad1faf0d8fc3504c5"
WEEK = "fixtures/week-2026-08-24"
TOLERANCE = D("1")
MINUS = "−"

# Minimum mapping (PRD v5.2-A R68), the same rows as packages/engine/src/fs-map.ts (the test compares them).
MAP = [
    ("PL", 12, "Revenue", "revenue", "Σ Revenue - LB Main Product, LB By Product, Trading (jurnal + penyesuaian tutup buku)"),
    ("PL", 13, "Cost of Sales", "cogs", "Σ COGS dan Other COGS; mode memo (O24 b): tanpa Selisih costing dan Biaya kapasitas idle"),
    ("PL", 15, "GROSS PROFIT", "gp", "Revenue + COGS (mode memo: tanpa baris v5)"),
    ("PL", 37, "TOTAL OPERATING EXPENSES", "opex", "semua baris selain Revenue, COGS, dan pendapatan/beban lain (termasuk penyusutan dan pajak)"),
    ("PL", 49, "TOTAL OTHER INCOME (EXPENSE)", "other", "Interest Expense, Other Income, Other Expense"),
    ("PL", 51, "INCOME (LOSS) BEFORE TAX", "nibt", "NIBT jurnal; mode memo: tanpa NIBT manajemen"),
    ("BS", 14, "Cash and cash equivalents", "kas", "saldo kas dan bank akhir bulan (TB awal + jurnal)"),
    ("BS", 17, "Trade Receivables - Third Parties", "piutang_pihak_ketiga", None),
    ("BS", 18, "Trade Receivables - Related Parties", "piutang_pihak_berelasi", None),
    ("BS", 23, "Inventory - Product", "persediaan_produk", "saldo buku besar (bs_line Inventory - Product)"),
    ("BS", 24, "Inventory - Trading", "persediaan_trading", "saldo buku besar (bs_line Inventory - Trading)"),
    ("BS", 47, "Trade Payables - Third Parties", "utang_dagang_pihak_ketiga", None),
    ("BS", 48, "Trade Payables - Affiliates", "utang_dagang_afiliasi", None),
    ("BS", 51, "Other Payable - Related Parties", "utang_lain_afiliasi", "pinjaman dan utang lain afiliasi"),
    ("BS", 77, "TOTAL EQUITY", "ekuitas", "Shareholder Equity + Retained earnings + laba tahun berjalan"),
]
APP_LABEL = {
    "revenue": "Revenue (jurnal)", "cogs": "COGS (jurnal)", "gp": "Gross profit (jurnal)", "opex": "Beban operasi (jurnal)",
    "other": "Pendapatan dan beban lain (jurnal)", "nibt": "NIBT (jurnal)", "kas": "Kas dan setara kas",
    "piutang_pihak_ketiga": "Piutang usaha pihak ketiga", "piutang_pihak_berelasi": "Piutang usaha pihak berelasi",
    "persediaan_produk": "Persediaan produk", "persediaan_trading": "Persediaan trading",
    "utang_dagang_pihak_ketiga": "Utang dagang pihak ketiga", "utang_dagang_afiliasi": "Utang dagang afiliasi",
    "utang_lain_afiliasi": "Utang lain pihak berelasi", "ekuitas": "Total ekuitas",
}
BS_SOURCE = [  # app_line, bs_line, side (+1 debit-positive asset, −1 credit-positive liability)
    ("kas", "Cash and cash equivalents", 1), ("piutang_pihak_ketiga", "Trade Receivables - Third Parties", 1),
    ("piutang_pihak_berelasi", "Trade Receivables - Related Parties", 1), ("persediaan_produk", "Inventory - Product", 1),
    ("persediaan_trading", "Inventory - Trading", 1), ("utang_dagang_pihak_ketiga", "Trade Payables - Third Parties", -1),
    ("utang_dagang_afiliasi", "Trade Payables - Affiliates", -1), ("utang_lain_afiliasi", "Other Payable - Related Parties", -1),
]
OTHER_INCOME = {"Interest Expense", "Other Income", "Other Expense"}


# ---------------------------------------------------------------------------------------------------- helpers
def read_json(rel):
    with open(os.path.join(ROOT, rel), encoding="utf-8") as f:
        return json.load(f)


def sha256(rel):
    h = hashlib.sha256()
    with open(os.path.join(ROOT, rel), "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def q2(v):
    """numeric(20,2) as fs_statement_line stores it (round half up)."""
    return D(v).quantize(D("0.01"), rounding=ROUND_HALF_UP)


def s2(v):
    q = q2(v)
    return format(D("0.00") if q == 0 else q, "f")  # never "-0.00"


def dump(data):
    return json.dumps(data, ensure_ascii=False, indent=1) + "\n"


def minus(s):
    return s.replace("-", MINUS, 1) if s.startswith("-") else s


def grp(n, dp):
    """1.234.567,89 (id-ID), no sign."""
    t = format(abs(n).quantize(D(1).scaleb(-dp), rounding=ROUND_HALF_UP), "f")
    i, _, f = t.partition(".")
    i = "{:,}".format(int(i)).replace(",", ".")
    return i + ("," + f if dp else "")


def jt(v):
    """Rp 8.361,8 jt / Rp −26,5 jt."""
    n = D(v) / D(1_000_000)
    return "Rp " + (MINUS if n.quantize(D("0.1"), rounding=ROUND_HALF_UP) < 0 else "") + grp(n, 1) + " jt"


def rp2(v):
    """Rp 1.234,56 / Rp −2,71 (exact rupiah, two decimals)."""
    n = D(v)
    return "Rp " + (MINUS if q2(n) < 0 else "") + grp(n, 2)


def sel(v):
    """Selisih: exact Rp below Rp 1 jt, else jt with the exact Rp beside it."""
    n = D(v)
    return rp2(n) if abs(n) < D(1_000_000) else f"{jt(n)} ({rp2(n)})"


# ---------------------------------------------------------------------------------------------------- FS side
def fs_rows(fs, col, statements):
    """Rows of the Management Report with a number in the month column, as fs_statement_line would hold them."""
    out = []
    for st, key, sheet in (("PL", "pl", "PL - Management Report"), ("BS", "bs", "BS - Management Report")):
        if st not in statements:
            continue
        for r in fs[key]:
            v = r["values"].get(col)
            if v is not None:
                out.append({"statement": st, "sheet": sheet, "row_no": r["row"], "line": r["line"], "amount": s2(v)})
    return out


def check_workbook(fs_rows_by_col):
    """Every FS row used is read from the workbook again: label and value must agree with fixtures/fs-2026-09 (inputs and cell golden)."""
    wb = openpyxl.load_workbook(os.path.join(ROOT, FS_XLSX), data_only=True)
    cells = read_json("fixtures/fs-2026-09/golden.json")
    n = 0
    for col, rows in fs_rows_by_col:
        for r in rows:
            ws = wb[r["sheet"]]
            hdr = {c: ws.cell(10, c).value for c in range(4, 13)}
            c = next(c for c, v in hdr.items() if isinstance(v, (dt.date, dt.datetime)) and v.date().isoformat() == col)
            label = str(ws.cell(r["row_no"], 2).value).strip()
            v = ws.cell(r["row_no"], c).value
            assert label == r["line"], f"{r['sheet']} b.{r['row_no']}: label {label!r} != {r['line']!r}"
            assert isinstance(v, (int, float)) and q2(repr(v)) == D(r["amount"]), f"{r['sheet']} b.{r['row_no']} {col}: {v!r} != {r['amount']}"
            ref = f"{r['sheet']}!{openpyxl.utils.get_column_letter(c)}{r['row_no']}"
            assert q2(cells[ref]) == D(r["amount"]), f"golden fs-2026-09 {ref}: {cells[ref]} != {r['amount']}"
            n += 1
    return n


# ---------------------------------------------------------------------------------------------------- app side
def app_pl(gl, coa, frm, to):
    """P&L per pl_line from the journal of [frm, to] (revenue +, cost −), as actualRowsV4 builds it."""
    line = {c["account"]: c["pl_line"] for c in coa if c["pl_line"]}
    out = {}
    for g in gl:
        if not g["date"] or not g["account"] or g["date"] < frm or g["date"] > to or g["account"] not in line:
            continue
        out[line[g["account"]]] = out.get(line[g["account"]], D(0)) + D(g["credit"] or 0) - D(g["debit"] or 0)
    return out


def fold_pl(by_line):
    is_rev = lambda l: l.startswith("Revenue")
    is_cogs = lambda l: l.startswith("COGS") or l.startswith("Other COGS")
    t = lambda p: sum((v for l, v in by_line.items() if p(l)), D(0))
    revenue, cogs, other = t(is_rev), t(is_cogs), t(lambda l: l in OTHER_INCOME)
    opex = t(lambda l: not is_rev(l) and not is_cogs(l) and l not in OTHER_INCOME)
    gp = revenue + cogs
    return {"revenue": revenue, "cogs": cogs, "gp": gp, "opex": opex, "other": other, "nibt": gp + opex + other}


def app_bs(tb, gl, coa, date, tb_date):
    """Trial balance + journal in (tb_date, date] per bs_line; equity = Shareholder Equity + Retained earnings + year profit."""
    c = {x["account"]: x for x in coa}
    bs, profit = {}, D(0)

    def add(acc, v):
        nonlocal profit
        x = c.get(acc)
        if x and x["type"] == "LabaRugi":
            profit += v
        elif x and x["type"] == "Neraca" and x["bs_line"]:
            bs[x["bs_line"]] = bs.get(x["bs_line"], D(0)) + v
        elif not (x and x["type"] == "Header"):
            assert v == 0, f"akun {acc} tidak bisa ditempatkan: {v}"

    for t in tb:
        add(str(t["account"]).strip(), D(t["balance"]))
    for g in gl:
        if g["date"] and g["account"] and tb_date < g["date"] <= date:
            add(str(g["account"]).strip(), D(g["debit"] or 0) - D(g["credit"] or 0))
    out = {a: bs.get(b, D(0)) * side for a, b, side in BS_SOURCE}
    out["ekuitas"] = -(bs.get("Shareholder Equity", D(0)) + bs.get("Retained earnings", D(0))) - profit
    parts = {"shareholder": -bs.get("Shareholder Equity", D(0)), "retained": -bs.get("Retained earnings", D(0)), "profit": -profit}
    return out, parts


# ---------------------------------------------------------------------------------------------------- reconciliation (oracle)
def recon_rows(app, fs, statements, partial, known):
    """Golden rows of the mapped pairs: selisih = app − FS, |selisih| ≤ Rp 1 = sama; cause per the file header."""
    by_fs = {(r["statement"], r["row_no"]): r for r in fs}
    rows = []
    for st, row, label, app_line, note in MAP:
        if st not in statements:
            continue
        f = by_fs.get((st, row))
        a = app.get(app_line)
        fsv = D(f["amount"]) if f else None
        sd = a - fsv if a is not None and fsv is not None else None
        status = "tidak ada di app" if a is None else "tidak ada di FS" if fsv is None else "sama" if abs(sd) <= TOLERANCE else "beda"
        cause = None if status == "sama" else "data belum ada" if status != "beda" or partial else "selisih"
        rows.append({
            "key": f"{st} b.{row}", "fs_line": label, "app_line": app_line,
            "app": s2(a) if a is not None else None, "fs": f["amount"] if f else None, "selisih": s2(sd) if sd is not None else None,
            "status": status, "cause": cause, "explain": known.get(app_line, "") if status != "sama" else "",
        })
    return rows


def summary(rows, fs, statements):
    mapped_keys = {(st, row) for st, row, *_ in MAP if st in statements}
    return {
        "mapped": len(rows), "sama": sum(r["status"] == "sama" for r in rows), "beda": sum(r["status"] == "beda" for r in rows),
        "missing": sum(r["status"].startswith("tidak ada") for r in rows),
        "unmapped": len([f for f in fs if (f["statement"], f["row_no"]) not in mapped_keys]),
    }


def c59(s):
    return "GAGAL" if s["beda"] else "BELUM DIISI" if s["missing"] or not s["mapped"] else "LULUS"


# ---------------------------------------------------------------------------------------------------- build
def build():
    fs = read_json("fixtures/fs-2026-09/inputs.json")
    week = read_json(f"{WEEK}/week.json")
    gl = read_json(f"{WEEK}/gl_lines.json")
    coa = read_json(f"{WEEK}/coa_map.json")
    tb = read_json(f"{WEEK}/tb_opening.json")
    month = read_json("fixtures/month-2026-08/inputs.json")
    tb_date = week["tb_opening_date"]
    through = max(g["date"] for g in gl if g["date"])
    assert through == month["cut_off"], (through, month["cut_off"])

    # --- Agustus: akhir bulan
    pl_aug = app_pl(gl, coa, "2026-08-01", "2026-08-31")
    pl_fold = fold_pl(pl_aug)
    anchor = read_json("fixtures/month-2026-08/golden-v4.json")["actual"]  # jt Rp, 1 desimal, dari O_PnL kolom MTD (independen dari skrip ini)
    j1 = lambda v: str((v / D(1_000_000)).quantize(D("0.1"), rounding=ROUND_HALF_UP))
    assert (j1(pl_fold["revenue"]), j1(pl_fold["cogs"]), j1(pl_fold["gp"]), j1(pl_fold["opex"] + pl_fold["other"]), j1(pl_fold["nibt"])) == \
        (anchor["revenue"], anchor["cogs"], anchor["gp"], anchor["opex"], anchor["nibt"]), "laba rugi app Agustus beda dari golden month-2026-08"
    bs_aug, eq_aug = app_bs(tb, gl, coa, "2026-08-31", tb_date)
    app_aug = {**pl_fold, **bs_aug}
    fs_aug = fs_rows(fs, "2026-08-31", ("PL", "BS"))
    by = {(r["statement"], r["row_no"]): D(r["amount"]) for r in fs_aug}
    missing_pl = [(l, by[("PL", row)]) for l, row in (("Salaries and allowance", 23), ("Depreciation & amortisation", 34), ("Interest Expense", 43)) if l not in pl_aug]
    miss_txt = ", ".join(f"{l} ({jt(v)})" for l, v in missing_pl)
    base = f"jurnal fixture baru sampai {dt.date.fromisoformat(through).day} Agu 2026; transaksi 27–31 Agu dan jurnal penutup bulan belum ada"
    known_aug = {
        "revenue": base, "cogs": base, "gp": base, "nibt": base,
        "opex": base + f"; baris yang FS punya dan jurnal fixture belum punya sama sekali: {miss_txt}",
        "other": base + "; Interest Expense belum dijurnal (FS Agu: " + jt(by[("PL", 43)]) + ")",
        "kas": base, "piutang_pihak_ketiga": base, "persediaan_produk": base, "persediaan_trading": base,
        "utang_dagang_pihak_ketiga": base + (f"; saldo utang dagang app bersaldo debit ({jt(bs_aug['utang_dagang_pihak_ketiga'])}): pembayaran lebih besar dari tagihan yang sudah dijurnal" if bs_aug["utang_dagang_pihak_ketiga"] < 0 else ""),
        "utang_lain_afiliasi": base, "ekuitas": base + "; laba tahun berjalan app memuat jurnal Agu s.d. 26 Agu saja, dan b.71 FS beda Rp 2,71 dari TB (lihat saldo awal)",
    }
    rows_aug = recon_rows(app_aug, fs_aug, ("PL", "BS"), True, known_aug)
    s_aug = summary(rows_aug, fs_aug, ("PL", "BS"))

    # --- Agustus: saldo awal 31 Jul (data lengkap: TB awal; FS Juli dipakai hanya di fixture, B-R91 mengimpor Agu dan Sep)
    bs_jul, eq_jul = app_bs(tb, gl, coa, tb_date, tb_date)
    fs_jul = fs_rows(fs, "2026-07-31", ("BS",))
    fsv = {r["row_no"]: D(r["amount"]) for r in fs_jul}
    re_fs, prof_fs = fsv[71], fsv[72]
    shareholder_fs = fsv[68]
    d_sh = eq_jul["shareholder"] - shareholder_fs
    d_re = eq_jul["retained"] - re_fs  # both credit-positive: a debit (deficit) balance is negative on each side
    d_pr = eq_jul["profit"] - prof_fs
    assert d_sh == 0 and abs(d_pr) < D("0.01") and q2(d_re) == q2(app_jul_eq_gap := (bs_jul["ekuitas"] - fsv[77])), "selisih ekuitas Juli tidak sepenuhnya dari b.71"
    known_jul = {
        "ekuitas": (f"selisih seluruhnya dari saldo laba ditahan awal: TB Accurate {rp2(eq_jul['retained'])}, BS - Management Report b.71 {rp2(re_fs)}, "
                    f"selisih {rp2(d_re)}; modal disetor b.68 dan laba tahun berjalan b.72 sama (selisih {rp2(d_sh)} dan {rp2(d_pr)}). "
                    f"Neraca FS sendiri tidak seimbang: total aset b.41 {rp2(fsv[41])} lawan liabilitas + ekuitas b.79 {rp2(fsv[79])}, selisih {rp2(fsv[41] - fsv[79])}"),
    }
    rows_jul = recon_rows(bs_jul, fs_jul, ("BS",), False, known_jul)
    s_jul = summary(rows_jul, fs_jul, ("BS",))

    # --- September: tidak ada jurnal di fixture
    fs_sep = fs_rows(fs, "2026-09-30", ("PL", "BS"))
    sep_note = "jurnal September tidak ada di fixture (jurnal fixture berhenti 26 Agu 2026); app belum punya angka September"
    rows_sep = recon_rows({}, fs_sep, ("PL", "BS"), True, {m[3]: sep_note for m in MAP})
    s_sep = summary(rows_sep, fs_sep, ("PL", "BS"))

    check_n = check_workbook([("2026-08-31", fs_aug), ("2026-07-31", fs_jul), ("2026-09-30", fs_sep)])

    map_json = [{"statement": st, "row_no": row, "fs_line": label, "app_line": a, "sign": 1, "note": note} for st, row, label, a, note in MAP]
    tol = "1"

    def app_json(app, lines):
        return [{"app_line": k, "amount": s2(app[k])} for k in lines if k in app]

    pl_lines = [{"pl_line": k, "amount": s2(v)} for k, v in sorted(pl_aug.items())]
    inp08 = {
        "month": "2026-08-01", "tolerance": tol, "costing_rows_mode": "memo", "map": map_json,
        "recons": [
            {
                "id": "akhir-bulan", "title": "Akhir bulan: laba rugi Agustus dan neraca 31 Agustus 2026", "fs_column": "2026-08-31", "statements": ["PL", "BS"],
                "app_data": {"week_fixture": "week-2026-08-24", "tb_opening_date": tb_date, "from": "2026-08-01", "journal_through": through, "month_end": "2026-08-31"},
                "pl_lines": pl_lines, "app": app_json(app_aug, [m[3] for m in MAP]), "fs": fs_aug,
            },
            {
                "id": "saldo-awal", "title": "Saldo awal: neraca 31 Juli 2026 (TB Accurate) lawan FS Juli", "fs_column": "2026-07-31", "statements": ["BS"],
                "app_data": {"week_fixture": "week-2026-08-24", "tb_opening_date": tb_date, "from": None, "journal_through": tb_date, "month_end": tb_date},
                "pl_lines": [], "app": app_json(bs_jul, [m[3] for m in MAP if m[0] == "BS"]), "fs": fs_jul,
            },
        ],
    }
    inp09 = {
        "month": "2026-09-01", "tolerance": tol, "costing_rows_mode": "memo", "map": map_json,
        "recons": [{
            "id": "akhir-bulan", "title": "Akhir bulan: laba rugi September dan neraca 30 September 2026", "fs_column": "2026-09-30", "statements": ["PL", "BS"],
            "app_data": None, "app_data_note": sep_note, "pl_lines": [], "app": [], "fs": fs_sep,
        }],
    }

    def gold(rows, s, rid):
        return {"id": rid, "rows": rows, "summary": s, "c59": c59(s)}

    gold08 = {
        "status": "diharapkan, dihitung dari data fixture (bukan angka disesuaikan); selisih 'data belum ada' = data Agustus di fixture belum lengkap",
        "unit": "Rp", "tolerance": tol,
        "recons": [gold(rows_aug, s_aug, "akhir-bulan"), gold(rows_jul, s_jul, "saldo-awal")],
    }
    gold09 = {
        "status": "diharapkan: tidak ada jurnal September di fixture, semua baris yang dipetakan 'tidak ada di app' dengan sebab 'data belum ada'",
        "unit": "Rp", "tolerance": tol, "recons": [gold(rows_sep, s_sep, "akhir-bulan")],
    }

    def source(name, extra):
        files = [
            (FS_XLSX, "PL - Management Report dan BS - Management Report, kolom Jul/Agu/Sep (label dan nilai dibaca ulang oleh skrip)"),
            ("fixtures/fs-2026-09/inputs.json", "baris FS per bulan (pl[].row, bs[].row)"),
            ("fixtures/fs-2026-09/golden.json", "sel-sel FS (sheet!sel) sebagai pembanding nilai"),
        ] + extra
        out = {"fixture": name, "generated_by": "scripts/build_fs_recon.py", "sources": []}
        for rel, used in files:
            e = {"file": rel, "sha256": sha256(rel), "used": used}
            if rel == FS_XLSX:
                assert e["sha256"] == FS_SHA, "workbook referensi berubah"
                e["sheets"] = ["PL - Management Report", "BS - Management Report"]
            out["sources"].append(e)
        return out

    app_srcs = [
        (f"{WEEK}/gl_lines.json", "jurnal 1–26 Agu 2026 (sisi app: laba rugi dan neraca)"),
        (f"{WEEK}/coa_map.json", "pl_line dan bs_line per akun"),
        (f"{WEEK}/tb_opening.json", "TB 31 Jul 2026 (saldo awal, O27)"),
        (f"{WEEK}/week.json", "tanggal TB awal"),
        ("fixtures/month-2026-08/inputs.json", "cut-off 26 Agu (batas jurnal fixture)"),
        ("fixtures/month-2026-08/golden-v4.json", "pembanding independen laba rugi Agustus s.d. 26 Agu (jt Rp)"),
    ]
    src08 = source("fs-recon-2026-08", app_srcs)
    src09 = source("fs-recon-2026-09", [(f"{WEEK}/gl_lines.json", "tanggal jurnal terakhir fixture (26 Agu 2026): tidak ada jurnal September")])
    return {
        "fixtures/fs-recon-2026-08": {"inputs.json": inp08, "golden.json": gold08, "source.json": src08},
        "fixtures/fs-recon-2026-09": {"inputs.json": inp09, "golden.json": gold09, "source.json": src09},
    }, {"aug": (rows_aug, s_aug, eq_aug, pl_aug, through), "jul": (rows_jul, s_jul, eq_jul), "sep": (rows_sep, s_sep), "n_checked": check_n, "tb_date": tb_date}


# ---------------------------------------------------------------------------------------------------- docs
def table(rows):
    out = ["| Baris FS | Baris app | App | FS | Selisih (app − FS) | Status | Sebab |", "|---|---|---:|---:|---:|---|---|"]
    for r in rows:
        why = "" if r["cause"] is None else r["cause"] + (f": {r['explain']}" if r["explain"] else "")
        out.append("| {k} {l} | {a} | {app} | {fs} | {sel} | {st} | {why} |".format(
            k=r["key"], l=r["fs_line"], a=APP_LABEL[r["app_line"]], app=jt(r["app"]) if r["app"] is not None else "belum ada",
            fs=jt(r["fs"]) if r["fs"] is not None else "belum ada", sel=sel(r["selisih"]) if r["selisih"] is not None else "—", st=r["status"], why=why))
    return "\n".join(out)


def doc(files, info):
    rows_aug, s_aug, eq_aug, pl_aug, through = info["aug"]
    rows_jul, s_jul, eq_jul = info["jul"]
    rows_sep, s_sep = info["sep"]
    fs = read_json("fixtures/fs-2026-09/inputs.json")
    by = {(st, r["row"]): r for st, k in (("PL", "pl"), ("BS", "bs")) for r in fs[k]}
    nibt_sep = D(by[("PL", 51)]["values"]["2026-09-30"])
    eq_sep = D(by[("BS", 77)]["values"]["2026-09-30"])
    L = []
    L += [
        "# Rekonsiliasi fixture app ↔ FS (R68, T15, K62)", "",
        "Dihasilkan oleh `python3 -I scripts/build_fs_recon.py` dari `fixtures/fs-recon-2026-08/golden.json` dan `fixtures/fs-recon-2026-09/golden.json` "
        "(`--check` memastikan berkas di repo sama dengan hasil skrip). Tes engine `packages/engine/test/fs-recon.test.ts` menjalankan `fsRecon` pada "
        "fixture yang sama dan harus menghasilkan tabel ini baris demi baris. Angka tidak disesuaikan agar lulus: bila data fixture belum lengkap, "
        "selisih yang diharapkan ditulis di sini dengan sebabnya.", "",
        "Format: `Rp 8.361,8 jt`, minus U+2212 (`Rp −26,5 jt`); selisih di bawah Rp 1 jt ditulis dalam rupiah penuh. Selisih = app − FS; |selisih| ≤ Rp 1 = `sama` (O26). "
        "Sebab: `data belum ada` (satu sisi tidak ada atau belum lengkap), `definisi` (arti baris beda, REVIEW-v5.1 §7), `selisih` (definisi sama, data lengkap: Finance harus menjelaskan).", "",
        "## 1. Sumber", "",
        f"- FS: `{FS_XLSX}` (sha256 `{FS_SHA}`), sheet `PL - Management Report` dan `BS - Management Report`; kolom Juli = J, Agustus = K, September = L. "
        f"Skrip membaca ulang label dan nilai {info['n_checked']} baris yang dipakai langsung dari workbook; semuanya sama dengan `fixtures/fs-2026-09/inputs.json`.",
        "- Salinan workbook acuan adalah yang ada di repo ini. Analisis lain memakai salinan berbeda (NIBT September −953,9 jt); di repo: NIBT September "
        f"{rp2(nibt_sep)}, ekuitas September {rp2(eq_sep)}.",
        f"- Sisi app: `fixtures/week-2026-08-24` (jurnal 1–{dt.date.fromisoformat(through).day} Agu 2026 dan TB 31 Jul 2026) dan `coa_map`. Jurnal September tidak ada di fixture.",
        "- Mode costing: produksi `costing_rows_mode='memo'` (O24 b): COGS untuk perbandingan FS tanpa baris v5 \"Selisih costing\" dan \"Biaya kapasitas idle\" "
        "(fixture Agustus tidak memuat baris itu; perilaku kedua mode diuji di `fs-recon.test.ts`).", "",
        "## 2. Pemetaan minimum (15 baris)", "",
        "| Baris FS | Label FS (diverifikasi) | Baris app | Tanda | Catatan |", "|---|---|---|---:|---|",
    ]
    for st, row, label, app_line, note in MAP:
        L.append(f"| {st} b.{row} | {label} | `{app_line}` ({APP_LABEL[app_line]}) | +1 | {note or ''} |")
    L += ["", "## 3. Agustus 2026: akhir bulan", "",
          f"Hasil: {s_aug['sama']} sama, {s_aug['beda']} beda, {s_aug['missing']} tidak ada; {s_aug['unmapped']} baris FS di luar pemetaan. **C59: {c59(s_aug)}** (diharapkan; bukan kegagalan engine).", "",
          table(rows_aug), "",
          f"Penjelasan: jurnal fixture berhenti {dt.date.fromisoformat(through).day} Agu, jadi semua selisih Agustus diberi sebab `data belum ada`. "
          "Bagian selisih yang berasal dari definisi atau selisih sungguhan baru bisa dipisahkan setelah jurnal 27–31 Agu dan jurnal penutup masuk (T20 di PRD B). "
          "Angka app s.d. 26 Agu cocok dengan golden independen `fixtures/month-2026-08/golden-v4.json` (revenue 6385,9; COGS −6038,5; GP 347,4; beban operasi dan lain −182,2; NIBT 165,2 jt) "
          "dan neraca v1 (`computeWeek`) pada 26 Agu (dicek di tes).", "",
          "## 4. Agustus 2026: saldo awal 31 Juli (data lengkap)", "",
          "Saldo awal = TB 31 Jul (O27). Ini satu-satunya perbandingan di fixture dengan data lengkap di kedua sisi. FS Juli hanya dipakai di fixture ini; B-R91 mengimpor kolom Agustus dan September.", "",
          f"Hasil: {s_jul['sama']} sama, {s_jul['beda']} beda, {s_jul['missing']} tidak ada. **C59: {c59(s_jul)}**.", "",
          table(rows_jul), "",
          "## 5. September 2026", "",
          f"Hasil: 0 sama, 0 beda, {s_sep['missing']} tidak ada di app; {s_sep['unmapped']} baris FS di luar pemetaan. **C59: {c59(s_sep)}** (bukan LULUS: baris yang tidak punya angka tidak pernah dibaca 0).", "",
          table(rows_sep), "",
          "## 6. Yang masih harus diputuskan Finance", "",
          f"1. **Ekuitas (b.77) berselisih {rp2(rows_jul[-1]['selisih'])} pada saldo awal, di atas toleransi Rp 1, walau data lengkap.** "
          "Saldo laba ditahan awal di `BS - Management Report` (b.71) sama di Juli, Agustus, dan September (−2.698.434.665,91), sedangkan TB Accurate dan sheet `BS` "
          "di workbook yang sama menunjukkan −2.698.434.668,62 pada ketiga bulan. Selama b.71 tidak dikoreksi, C59 pada b.77 GAGAL setiap bulan, juga setelah semua jurnal masuk. "
          "Pilihan Finance: koreksi b.71 di workbook, atau override C59 (policy B-R86). Engine tidak menutup selisih ini (K62).",
          "2. Baris FS di luar pemetaan minimum (rincian beban, aset tetap, pajak dibayar di muka, akrual) berstatus `tidak ada di app` dengan sebab `definisi` sampai O27 (b) dikerjakan (per akun TB).",
          "3. Bukti `LULUS` pada data lengkap Agustus dan September baru ada setelah jurnal lengkap dan FINAL di KGR-OS (T20 di PRD B).", ""]
    return "\n".join(L)


def main():
    files, info = build()
    out = {}
    for d, fs_ in files.items():
        for name, data in fs_.items():
            out[os.path.join(d, name)] = dump(data)
    out["docs/v5.2/fs-recon-fixture.md"] = doc(files, info)
    check = "--check" in sys.argv
    bad = 0
    for rel, text in out.items():
        path = os.path.join(ROOT, rel)
        if check:
            cur = open(path, encoding="utf-8").read() if os.path.exists(path) else None
            if cur != text:
                print(f"BEDA {rel}")
                bad += 1
            continue
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"ditulis {rel}")
    if check:
        print("fs-recon: berkas di repo = hasil skrip" if not bad else f"fs-recon: {bad} berkas beda dari hasil skrip")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
