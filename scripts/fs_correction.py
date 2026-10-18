#!/usr/bin/env python3
"""FS correction (K70, Q-v5.2-66: a correction is a new version) from a newer copy of the FS workbook.

  python3 -I scripts/fs_correction.py --source "<newer FS v2.xlsx>" [--out docs/v5.3/kgr-os-data-fs-v2.sql]

Reads "PL - Management Report" and "BS - Management Report" of August and September 2026 from the newer copy (same layout
as scripts/fs_import.py: label in column B, month-end column found in header row 10) and from the copy in the repo
(reference/…, the source of version 1 in KGR-OS, docs/v5.2/kgr-os-data-fs.sql). Every row whose amount differs becomes
version 2 of that row (locked, source file and sha256 of the newer copy); unchanged rows keep version 1, and the app reads
the newest version of each row (lib/v5/fs-recon-data.ts). A row label that changed, or a row present in only one copy,
stops the script (that is a layout change, not a correction). Nothing is adjusted (K62). The SQL refuses to run unless
version 1 of the old copy is there, inserts nothing twice, and ends with a check table whose every row must be true.
"""
import argparse, datetime as dt, hashlib, os, sys
from decimal import ROUND_HALF_UP, Decimal as D
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OLD = os.path.join(ROOT, "reference", "09._Laporan_Keuangan_Per_Sep_2026_PT_KGR_v2.xlsx")
MONTHS = [("2026-08-01", "2026-08-31"), ("2026-09-01", "2026-09-30")]
SHEETS = [("PL", "PL - Management Report"), ("BS", "BS - Management Report")]


def q2(v):
    return D(repr(v) if isinstance(v, float) else str(v)).quantize(D("0.01"), rounding=ROUND_HALF_UP)


def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for c in iter(lambda: f.read(1 << 20), b""):
            h.update(c)
    return h.hexdigest()


def rows_of(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for month, end in MONTHS:
        for st, sheet in SHEETS:
            ws = wb[sheet]
            col = next(c for c in range(4, 13) if isinstance(ws.cell(10, c).value, (dt.date, dt.datetime)) and ws.cell(10, c).value.date().isoformat() == end)
            for r in range(11, ws.max_row + 1):
                label, v = ws.cell(r, 2).value, ws.cell(r, col).value
                if label is None or str(label).strip() == "" or isinstance(v, bool) or not isinstance(v, (int, float)):
                    continue
                out[(month, st, r)] = (sheet, str(label).strip(), q2(v))
    return out


def lit(s):
    return "'" + s.replace("'", "''") + "'"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", required=True)
    ap.add_argument("--out", default=os.path.join(ROOT, "docs", "v5.3", "kgr-os-data-fs-v2.sql"))
    a = ap.parse_args()
    old, new = rows_of(OLD), rows_of(a.source)
    if set(old) != set(new):
        sys.exit(f"baris berbeda antar salinan (perubahan tata letak, bukan koreksi): {sorted(set(old) ^ set(new))[:8]}")
    bad = [k for k in old if old[k][1] != new[k][1]]
    if bad:
        sys.exit(f"label berubah: {[(k, old[k][1], new[k][1]) for k in bad[:8]]}")
    changed = sorted(k for k in old if old[k][2] != new[k][2])
    if not changed:
        sys.exit("tidak ada baris yang berubah: koreksi tidak perlu")
    so, sn = sha(OLD), sha(a.source)
    name = os.path.basename(a.source)
    values = ",\n".join(f"  ({lit(m)}::date, {lit(st)}, {lit(new[(m, st, r)][0])}, {r}, {lit(new[(m, st, r)][1])}, {new[(m, st, r)][2]}, {old[(m, st, r)][2]})" for m, st, r in changed)
    checks = [
        f"  ('{len(changed)} baris versi 2 dari berkas baru, terkunci',\n   (select count(*) = {len(changed)} and count(locked_at) = {len(changed)} and bool_and(source_sha256 = {lit(sn)}) from public.fs_statement_line where version = 2 and month in ('2026-08-01','2026-09-01')))",
    ]
    for m, st, r in changed:
        checks.append(f"  ('{st} {m[:7]} b.{r} {new[(m, st, r)][1][:40]}: versi terbaru = {new[(m, st, r)][2]}',\n   (select amount = {new[(m, st, r)][2]} from public.fs_statement_line where month = {lit(m)} and statement = {lit(st)} and row_no = {r} order by version desc limit 1))")
    text = f"""-- KGR-OS: koreksi FS acuan (K70, Q-v5.2-66: koreksi = versi baru), dibangkitkan scripts/fs_correction.py (jangan disunting tangan).
-- Lama (versi 1, docs/v5.2/kgr-os-data-fs.sql): reference/09._Laporan_Keuangan_Per_Sep_2026_PT_KGR_v2.xlsx sha256 {so}
-- Baru (versi 2): {name} sha256 {sn} (salinan 18 Okt; tidak di-commit)
-- {len(changed)} baris berubah; baris lain tetap versi 1. App membaca versi terbaru tiap baris (lib/v5/fs-recon-data.ts).
-- Prasyarat: versi 1 dari berkas lama ada untuk Agustus dan September. Bisa dijalankan ulang. Tabel cek di akhir: semua true.

do $$
begin
  if (select count(*) from public.fs_statement_line where version = 1 and source_sha256 = {lit(so)} and month in ('2026-08-01','2026-09-01')) = 0 then
    raise exception 'versi 1 FS dari berkas lama tidak ada: jalankan dulu docs/v5.2/kgr-os-data-fs.sql';
  end if;
end $$;

insert into public.fs_statement_line (month, statement, sheet, row_no, line, amount, source_file, source_sha256, version, locked_at)
select v.month, v.statement, v.sheet, v.row_no, v.line, v.amount, {lit(name)}, {lit(sn)}, 2, now()
from (values
{values}
) as v(month, statement, sheet, row_no, line, amount, old_amount)
where exists (select 1 from public.fs_statement_line o where o.month = v.month and o.statement = v.statement and o.row_no = v.row_no and o.version = 1 and o.amount = v.old_amount)
on conflict (month, statement, row_no, version) do nothing;

select cek, ada from (values
{(',' + chr(10)).join(checks)}
) as t(cek, ada);
"""
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    open(a.out, "w", encoding="utf-8").write(text)
    print(f"{os.path.relpath(a.out, ROOT)}: {len(changed)} baris berubah")
    for m, st, r in changed:
        print(f"  {st} {m[:7]} b.{r} {new[(m, st, r)][1]}: {old[(m, st, r)][2]} -> {new[(m, st, r)][2]}")


if __name__ == "__main__":
    main()
