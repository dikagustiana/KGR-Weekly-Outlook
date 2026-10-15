"""Template kas harian for Finance (PRD v3 §4.7, PLAN Fase 5) → reference/templates/Template_Kas_Harian.xlsx.

Two sheets: "Saldo" (tanggal, no_rekening, saldo_awal, saldo_akhir, catatan) and "Rencana" (tanggal, no_rekening,
kategori_cf, deskripsi, jumlah, id_jadwal_tetap). kategori_cf has a dropdown of the cash-flow categories of the chart of
accounts (supabase/seed/04_coa_map.sql, column cf_category_as_counter). Run: python3 -I scripts/make_cash_template.py
"""
import os
import re

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def cf_categories():
    sql = open(os.path.join(ROOT, "supabase", "seed", "04_coa_map.sql"), encoding="utf-8").read()
    cols = re.search(r"insert into public\.coa_map\s*\(([^)]*)\)", sql, re.I).group(1).split(",")
    idx = [c.strip() for c in cols].index("cf_category_as_counter")
    cats = set()
    for row in re.findall(r"\(([^()]*(?:\([^()]*\)[^()]*)*)\)\s*[,;]", sql.split("values", 1)[1]):
        vals = re.findall(r"'((?:[^']|'')*)'|(null|true|false|-?\d+(?:\.\d+)?)", row, re.I)
        flat = [a.replace("''", "'") if a else (None if b.lower() == "null" else b) for a, b in vals]
        if len(flat) > idx and flat[idx]:
            cats.add(flat[idx])
    return sorted(cats)


def main():
    cats = cf_categories()
    wb = openpyxl.Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2F6DB5")

    s = wb.active
    s.title = "Saldo"
    s.append(["tanggal", "no_rekening", "saldo_awal", "saldo_akhir", "catatan"])
    s.append(["2026-10-13", "BCA-0000000000", 0, 0, "contoh — hapus baris ini"])
    for c, w in zip("ABCDE", (12, 20, 18, 18, 40)):
        s.column_dimensions[c].width = w

    r = wb.create_sheet("Rencana")
    r.append(["tanggal", "no_rekening", "kategori_cf", "deskripsi", "jumlah", "id_jadwal_tetap"])
    r.append(["2026-10-14", "BCA-0000000000", cats[0] if cats else "", "contoh: + masuk / − keluar — hapus baris ini", 0, ""])
    for c, w in zip("ABCDEF", (12, 20, 42, 44, 18, 38)):
        r.column_dimensions[c].width = w

    lists = wb.create_sheet("Daftar")
    lists.append(["kategori_cf"])
    for c in cats:
        lists.append([c])
    lists.sheet_state = "hidden"
    dv = DataValidation(type="list", formula1=f"=Daftar!$A$2:$A${len(cats) + 1}", allow_blank=True, showErrorMessage=True,
                        errorTitle="Kategori arus kas", error="Pilih kategori dari daftar (COA)")
    dv.add("C2:C2000")
    r.add_data_validation(dv)

    for ws in (s, r):
        for cell in ws[1]:
            cell.font, cell.fill = head, fill
        ws.freeze_panes = "A2"
    out = os.path.join(ROOT, "reference", "templates", "Template_Kas_Harian.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"{os.path.relpath(out, ROOT)}: {len(cats)} kategori arus kas")


if __name__ == "__main__":
    main()
