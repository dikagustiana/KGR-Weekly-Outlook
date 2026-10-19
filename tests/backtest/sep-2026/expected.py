"""Backtest September 2026 — expected PL/BS from the FS workbook IN THE REPO (named source, sha256 recorded).
Run from the repo root: python3 -I tests/backtest/sep-2026/expected.py
Cell keys name the real cell: '<sheet>'!<column><row>, September = column L, August = column K (row 10 = month end)."""
import hashlib, json, os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
# BACKTEST_FS / BACKTEST_OUT: another copy of the FS workbook (e.g. the 18 Okt version in BACKTEST_DIR, not committed) and its output
SRC = os.environ.get("BACKTEST_FS") or os.path.join(ROOT, "reference", "09._Laporan_Keuangan_Per_Sep_2026_PT_KGR_v2.xlsx")
OUT = os.environ.get("BACKTEST_OUT") or os.path.join(ROOT, "docs", "v5.2", "backtest-sep", "expected.json")
COLS = {"2026-08": ("K", 11), "2026-09": ("L", 12)}

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
sha = hashlib.sha256(open(SRC, "rb").read()).hexdigest()
out = {"source": {"file": os.path.relpath(SRC, ROOT) if SRC.startswith(ROOT) else os.path.basename(SRC), "sha256": sha,
                  "sheets": ["PL - Management Report", "BS - Management Report"]}, "months": {}}
for month, (letter, col) in COLS.items():
    m = {"pl": {}, "bs": {}}
    for key, sheet, rows in (("pl", "PL - Management Report", range(12, 60)), ("bs", "BS - Management Report", range(12, 80))):
        ws = wb[sheet]
        for r in rows:
            label, v = ws.cell(row=r, column=2).value, ws.cell(row=r, column=col).value
            if label is not None and isinstance(v, (int, float)):
                m[key][f"'{sheet}'!{letter}{r}"] = {"row": r, "label": str(label).strip(), "value": round(float(v), 2)}
    out["months"][month] = m
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(out, f, indent=1, ensure_ascii=False)
for month, (letter, _) in COLS.items():
    pl, bs = out["months"][month]["pl"], out["months"][month]["bs"]
    g = lambda d, s, r: d.get(f"'{s}'!{letter}{r}", {}).get("value")
    print(month, "revenue", g(pl, "PL - Management Report", 12), "GP", g(pl, "PL - Management Report", 15),
          "NIBT", g(pl, "PL - Management Report", 51), "kas", g(bs, "BS - Management Report", 14), "ekuitas", g(bs, "BS - Management Report", 77))
